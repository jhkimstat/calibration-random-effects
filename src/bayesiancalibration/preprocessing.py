"""Fixed loading-only preparation for the supplied synthetic calibration data.

The supplied synthetic rows use increasing depth. The offset matcher also
accepts irregular, repeated, and unsorted per-curve depths. Projected data
and coefficients are stacked site/run first, with five loading coefficients.
No evaluation truth enters the preparation or the library standardization.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import platform
import tempfile
from dataclasses import dataclass
from importlib.metadata import version
from pathlib import Path

import jax
import jax.numpy as jnp
import numpy as np

from bayesiancalibration.state import ThetaStandardization
from bayesiancalibration.data import project_loading_curves


_FILES = (
    "2026BKA.xlsx", "BKA5_2_depth.csv", "BKA5_2_load.csv",
    "syn_theta_spatial_depth.csv", "syn_theta_spatial_load.csv",
    "theta_spatial.csv",
)
_PARAMETER_NAMES = ("EA", "EM", "etr")


@dataclass(frozen=True)
class PreparedSyntheticData:
    """Frozen loading data; truth is for evaluation only, never model inputs."""

    theta_s_dagger: np.ndarray  # (r,3), physical library parameters
    theta_s_tilde: np.ndarray  # (r,3), standardized library inputs
    theta_bar_dagger: np.ndarray  # (3,)
    D_theta: np.ndarray  # (3,3)
    F_s: np.ndarray  # (r,5), fitted library coefficients
    y_tilde: np.ndarray  # (n*5,), site-major projected observations
    R: np.ndarray  # (n*5,n*5), site-major QR factors
    s: np.ndarray  # (n,2), field coordinates
    theta_f_truth_for_evaluation: np.ndarray  # (n,3), never fitted
    library_depth: np.ndarray  # (m_s,), after shift
    library_t: np.ndarray  # (m_s,), after normalization
    library_load_shifted: np.ndarray  # (r,m_s)
    field_depth: np.ndarray  # (m_f,)
    field_t: np.ndarray  # (m_f,)
    field_load_noisy: np.ndarray  # (n,m_f)
    field_noise: np.ndarray  # (n,m_f), exact realized N(0,1) draws
    slope_offsets: np.ndarray  # regular candidate offset grid
    slope_errors: np.ndarray  # absolute mean-slope difference; NaN if unestimable
    mean_library_slopes: np.ndarray  # (candidate offsets,)
    field_slopes: np.ndarray  # (n,)
    metadata: dict



def _load_csv(path: Path, *, header: tuple[str, ...] | None = None) -> np.ndarray:
    """Keep source row order and reject nonnumeric or malformed CSV records."""

    with path.open(encoding="utf-8-sig") as file:
        labels = tuple(part.strip() for part in next(csv.reader(file)))
    if header is not None and labels != header:
        raise ValueError(f"Unexpected columns in {path.name}: {labels}")
    values = np.loadtxt(path, delimiter=",", skiprows=1, ndmin=2, dtype=np.float64)
    if values.shape[1] != len(labels) or not np.all(np.isfinite(values)):
        raise ValueError(f"Invalid numeric array in {path.name}")
    return values


def _read_library_parameters(path: Path) -> np.ndarray:
    """Read the 20 fixed Case rows using the standard openpyxl XLSX reader."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected workbook sheets")
        rows = list(workbook["Sheet1"].values)
    finally:
        workbook.close()
    if (len(rows) != 21 or len(rows[0]) != 4
            or rows[0][0] != "Case"
            or not all(name in str(label) for name, label in zip(
                ("Ea", "Em", "etr"), rows[0][1:]))):
        raise ValueError("Unexpected library workbook layout")
    if any(row[0] != i for i, row in enumerate(rows[1:], start=1)):
        raise ValueError("Library Case IDs must follow curve rows 1–20")
    theta = np.asarray([row[1:] for row in rows[1:]], dtype=np.float64)
    if theta.shape != (20, 3) or not np.all(np.isfinite(theta)):
        raise ValueError("Invalid library calibration parameters")
    return theta


def prepare_synthetic_data(
    input_dir: str | Path, *, seed: int = 1024, candidate_step: float = 0.8
) -> PreparedSyntheticData:
    """Prepare all 20 library runs and 60 sites with one fixed noisy field draw."""

    if not jax.config.jax_enable_x64:
        raise RuntimeError("Enable jax_enable_x64 before preprocessing")
    if not isinstance(seed, int) or seed < 0 or seed >= 2**32:
        raise ValueError("seed must be an unsigned 32-bit integer")
    source = Path(input_dir)
    paths = {name: source / name for name in _FILES}
    theta_s_dagger = _read_library_parameters(paths["2026BKA.xlsx"])
    h_s_rows = _load_csv(paths["BKA5_2_depth.csv"])
    y_s = _load_csv(paths["BKA5_2_load.csv"])
    h_f_rows = _load_csv(paths["syn_theta_spatial_depth.csv"])
    y_f = _load_csv(paths["syn_theta_spatial_load.csv"])
    truth = _load_csv(paths["theta_spatial.csv"], header=_PARAMETER_NAMES)
    if (h_s_rows.shape != (20, 501) or y_s.shape != h_s_rows.shape
            or h_f_rows.shape != (60, 501) or y_f.shape != h_f_rows.shape
            or truth.shape != (60, 3)):
        raise ValueError("Unexpected synthetic curve or truth shapes")
    if (not np.allclose(h_s_rows, h_s_rows[0], rtol=0, atol=1e-10)
            or not np.allclose(h_f_rows, h_f_rows[0], rtol=0, atol=1e-10)):
        raise ValueError("All curves in each source must share a depth grid")
    h_s, h_f = h_s_rows[0], h_f_rows[0]
    if (not np.isclose(h_s[0], 0.0) or not np.isclose(h_f[0], 0.0)
            or not np.isclose(h_s[-1], 400.0) or not np.isclose(h_f[-1], 400.0)
            or np.any(np.diff(h_s) <= 0) or np.any(np.diff(h_f) <= 0)):
        raise ValueError("Unexpected synthetic depth endpoints or order")

    noise = np.asarray(jax.random.normal(
        jax.random.key(seed), y_f.shape, dtype=jnp.float64
    ))
    y_f_noisy = y_f + noise
    projection = project_loading_curves(
        h_s, y_s, h_f, y_f_noisy, candidate_step=candidate_step,
    )
    h0 = projection.h0
    offsets, errors = projection.slope_offsets, projection.slope_errors
    slopes_s, slopes_f = projection.mean_library_slopes, projection.field_slopes
    F_s, y_tilde, R = projection.F_s, projection.y_tilde, projection.R
    h_s_shifted, t_s = projection.library_depth, projection.library_t
    y_s_shifted, t_f = projection.library_load_shifted, projection.field_t
    coords = np.column_stack((np.tile(np.arange(0.0, 120.0, 6.0), 3),
                              np.repeat([12.0, 6.0, 0.0], 20)))

    standardization = ThetaStandardization.from_library(theta_s_dagger)
    theta_s_tilde = np.asarray(standardization.to_standardized(theta_s_dagger))
    selected_index = int(np.searchsorted(offsets, h0))
    selected_error = float(errors[selected_index])
    metadata = {
        "stage": "14a", "dataset": "supplied_synthetic_loading_only",
        "branch_sizes": [5], "stacking": "site-major, loading coefficients 1..5",
        "parameter_names": list(_PARAMETER_NAMES), "parameter_units": None,
        "depth_units": None, "load_units": None, "coordinate_units": None,
        "load_scaling": "unchanged numeric values",
        "field_noise": {"distribution": "normal", "mean": 0.0, "sd": 1.0,
                        "seed": seed, "key_impl": str(jax.random.key_impl(jax.random.key(seed))),
                        "added_before": ["slope_matching", "QR_projection"]},
        "offset_rule": {
            "window": 4.0, "slope_fit": "OLS with intercept on observed points in each window",
            "interpolation": "linear only for post-alignment library baseline",
            "search": "regular candidate offset grid",
            "candidate_step_max": candidate_step,
            "candidate_count": int(offsets.size),
            "candidate_domain": [float(offsets[0]), float(offsets[-1])],
            "tie_break": "smallest evaluated offset with exactly minimal error",
            "h0": h0,
            "mean_field_slope": float(slopes_f.mean()),
            "mean_library_slope_at_h0": float(slopes_s[selected_index]),
            "absolute_error_at_h0": selected_error,
            "minimum_evaluated_error": float(np.nanmin(errors)),
            "unestimable_candidate_count": int(np.sum(~np.isfinite(errors))),
        },
        "basis": "R splines::bs degree=3, knots=(1/3,2/3), boundaries=(0,1), intercept=FALSE",
        "library_coefficient_fit": "full-rank economic QR, no coefficient rescaling",
        "field_projection": "economic QR, positive diagonal R",
        "library_fit_rms": projection.library_fit_rms,
        "field_projection_discarded_rms": projection.field_projection_discarded_rms,
        "source_sha256": {name: hashlib.sha256(path.read_bytes()).hexdigest()
                          for name, path in paths.items()},
        "source_directory": str(source.resolve()),
        "software": {name: version(name) for name in (
            "bayesiancalibration", "jax", "jaxlib", "numpy", "scipy", "openpyxl"
        )},
        "python": platform.python_version(),
    }
    return PreparedSyntheticData(
        theta_s_dagger, theta_s_tilde, np.asarray(standardization.theta_bar_dagger),
        np.asarray(standardization.D_theta), F_s, y_tilde, R, coords, truth,
        h_s_shifted, t_s, y_s_shifted, h_f, t_f, y_f_noisy, noise,
        offsets, errors, slopes_s, slopes_f, metadata,
    )


def save_prepared(data: PreparedSyntheticData, output_dir: str | Path) -> None:
    """Write a pickle-free archive and a diagnostic plot for this fixed input."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    from dataclasses import fields

    # Matplotlib and fontconfig need a writable cache under sandboxed runs.
    cache = str(Path(tempfile.gettempdir()) / "bayesiancalibration-plot-cache")
    os.environ.setdefault("MPLCONFIGDIR", cache)
    os.environ.setdefault("XDG_CACHE_HOME", cache)
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    arrays = {field.name: getattr(data, field.name) for field in fields(data)
              if field.name != "metadata"}
    np.savez_compressed(destination / "synthetic_preprocessing.npz", **arrays,
                        metadata=np.asarray(json.dumps(data.metadata, sort_keys=True)))
    offsets, errors = data.slope_offsets, data.slope_errors
    h0 = data.metadata["offset_rule"]["h0"]
    chosen = data.metadata["offset_rule"]["absolute_error_at_h0"]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4), layout="constrained")
    for ax in axes:
        ax.plot(offsets, errors, ".", color="tab:blue", ms=2)
        ax.axvline(h0, color="tab:red", ls="--", lw=1.1)
        ax.plot([h0], [chosen], "o", color="tab:red", ms=4)
        ax.set(xlabel="Library offset h0", ylabel="Absolute mean slope error")
        ax.grid(alpha=0.2)
    axes[0].set_title("Full range: regular offset grid")
    local_min, local_max = max(offsets[0], h0 - 4.0), min(offsets[-1], h0 + 4.0)
    local_errors = errors[(offsets >= local_min) & (offsets <= local_max)]
    axes[1].set_xlim(local_min, local_max)
    axes[1].set_ylim(0.0, 1.1 * max(float(np.nanmax(local_errors)), chosen, 1e-12))
    axes[1].set_title(f"Selected h0 = {h0:.6g}; error = {chosen:.3g}")
    fig.savefig(destination / "slope_error.png", dpi=180)
    plt.close(fig)


def main() -> None:
    """CLI keeps source selection explicit and the realized noise reproducible."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--seed", type=int, default=1024)
    parser.add_argument("--candidate-step", type=float, default=0.8)
    arguments = parser.parse_args()
    jax.config.update("jax_enable_x64", True)
    prepared = prepare_synthetic_data(
        arguments.input_dir, seed=arguments.seed,
        candidate_step=arguments.candidate_step,
    )
    save_prepared(prepared, arguments.output_dir)
    print(json.dumps({"h0": prepared.metadata["offset_rule"]["h0"],
                      "absolute_slope_error": prepared.metadata["offset_rule"][
                          "absolute_error_at_h0"],
                      "library_shape": list(prepared.F_s.shape),
                      "projected_shape": list(prepared.y_tilde.shape)}, indent=2))


if __name__ == "__main__":
    main()
